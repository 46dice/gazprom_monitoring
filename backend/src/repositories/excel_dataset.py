import re
import threading
import time
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile

from openpyxl import load_workbook

from schemas.dataset import CellValue, ColumnSchema, DataType, DatasetSchema

# Значения, которыми в файле обозначают «нет данных».
EMPTY_MARKERS = {"", "-", "—"}
# Доля значений нужного типа, чтобы считать столбец числовым/датой: в файле есть мусор.
TYPE_THRESHOLD = 0.9


class ExcelDatasetRepository:
    """Читает Excel и держит результат в памяти, пока файл не изменился."""

    def __init__(self, path: Path, header_rows: int):
        self._path = path
        self._header_rows = header_rows
        self._cache: DatasetSchema | None = None
        # Несколько запросов сразу после правки файла не должны разбирать его параллельно.
        self._lock = threading.Lock()

    def version(self) -> int:
        # Миллисекунды, а не наносекунды: наносекунды не влезают в Number в JS без потери точности.
        try:
            return self._path.stat().st_mtime_ns // 1_000_000
        except FileNotFoundError:
            return 0

    def get(self) -> DatasetSchema:
        with self._lock:
            version = self.version()
            if self._cache is None or self._cache.version != version:
                self._cache = self._read(version)
            return self._cache

    def _read(self, version: int) -> DatasetSchema:
        raw = self._load_rows()
        header_group, header_sub = raw[0], raw[1]
        data = [row for row in raw[self._header_rows:] if any(v not in (None, "") for v in row)]

        columns: list[ColumnSchema] = []
        indexes: list[int] = []
        group = ""
        for i, (top, sub) in enumerate(zip(header_group, header_sub)):
            top, sub = _header(top), _header(sub)
            # Объединённая ячейка группы заполнена только в первом столбце —
            # подстолбцы справа от неё наследуют имя группы.
            if top:
                group = top
            elif not sub:
                continue
            if not group:
                continue
            values = [_clean(row[i]) for row in data]
            columns.append(
                ColumnSchema(
                    key=f"{group} / {sub}" if sub and top != sub else group,
                    header_group=group,
                    header_sub=sub if sub != group else None,
                    data_type=_detect_type(values),
                )
            )
            indexes.append(i)

        rows = [[_to_json(_clean(row[i])) for i in indexes] for row in data]
        return DatasetSchema(version=version, file_name=self._path.name, columns=columns, rows=rows)

    def _load_rows(self, attempts: int = 5, delay_sec: float = 0.5) -> list[tuple]:
        # Пока Excel сохраняет файл, он может быть недописан или заблокирован.
        for attempt in range(attempts):
            try:
                wb = load_workbook(self._path, read_only=True, data_only=True)
                try:
                    return list(wb.worksheets[0].iter_rows(values_only=True))
                finally:
                    wb.close()
            except (PermissionError, BadZipFile, OSError):
                if attempt == attempts - 1:
                    raise
                time.sleep(delay_sec)
        return []


def _header(value) -> str | None:
    if value is None or value == 0:
        return None
    text = re.sub(r"\s+", " ", str(value)).strip()
    return text or None


def _clean(value) -> CellValue | datetime:
    if isinstance(value, str):
        value = value.strip()
        return None if value in EMPTY_MARKERS else value
    if isinstance(value, (datetime, date)):
        return value
    if isinstance(value, (int, float)):
        return value
    return None


def _detect_type(values: list) -> DataType:
    filled = [v for v in values if v is not None]
    if not filled:
        return "string"
    if sum(isinstance(v, (datetime, date)) for v in filled) / len(filled) >= TYPE_THRESHOLD:
        return "date"
    if sum(isinstance(v, (int, float)) for v in filled) / len(filled) >= TYPE_THRESHOLD:
        return "number"
    return "string"


def _to_json(value) -> CellValue:
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return value
